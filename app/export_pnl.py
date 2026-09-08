"""Excel-export: Top-N objecten als development-P&L (in de stijl van de
Hoogstraat-91-template). Eén 'Aannames'-tab, een 'Top 10'-overzicht en per
object een P&L-tab met echte formules — zodat het bestand herrekent en
aanpasbaar is.

Gebruik:
    from .export_pnl import build_pnl_workbook
    path = build_pnl_workbook("/Users/.../Desktop", region="grote_steden", n=10)
"""
from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .scenarios import (city_medians_all, get_profile, merged_params, top_listings)

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")            # invoer / aan te passen
BLACK = Font(name=FONT, color="000000")
GREEN = Font(name=FONT, color="008000")           # link naar andere tab
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14)
HEAD = Font(name=FONT, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3B57")
SUB_FILL = PatternFill("solid", fgColor="E8EEF4")
EUR = '€ #,##0'
PCT = '0.0%'
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(bottom=THIN)


def _safe_sheet_name(name: str, used: set[str]) -> str:
    n = re.sub(r'[\[\]:*?/\\]', " ", name)[:28].strip() or "Object"
    base, i = n, 2
    while n in used:
        n = f"{base[:25]} {i}"; i += 1
    used.add(n)
    return n


def _set(ws, cell, value, font=BLACK, fmt=None, align=None, fill=None, border=False):
    c = ws[cell]
    c.value = value
    c.font = font
    if fmt:
        c.number_format = fmt
    if align:
        c.alignment = Alignment(horizontal=align)
    if fill:
        c.fill = fill
    if border:
        c.border = BORDER
    return c


# ── Aannames-tab ──────────────────────────────────────────────────────
# (rij-nummers hier worden hard gerefereerd vanuit de object-tabs)
A = {
    "ovb": 4, "notaris": 5, "dd": 6, "taxatie": 7,
    "verbouw_m2": 8, "contingency": 9, "splitskosten": 10,
    "rente": 11, "looptijd": 12, "holding_overig": 13,
    "courtage": 14, "verkoop_vast": 15, "split_verlies": 16, "band": 17,
}


def _build_aannames(ws, p: dict):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 60
    _set(ws, "A1", "AANNAMES & BRONNEN — DealRadar model", TITLE)
    _set(ws, "A3", "Aanname", HEAD, fill=HEAD_FILL)
    _set(ws, "B3", "Waarde", HEAD, fill=HEAD_FILL, align="right")
    _set(ws, "C3", "Eenheid", HEAD, fill=HEAD_FILL)
    _set(ws, "D3", "Toelichting", HEAD, fill=HEAD_FILL)
    rows = [
        ("ovb", "Overdrachtsbelasting", p["ovb_pct"] / 100, "%", "8% (woning niet als hoofdverblijf), 2026"),
        ("notaris", "Notaris leveringsakte + Kadaster", p["kk_vast"], "€", "Onderdeel aankoopkosten (modelaanname)"),
        ("dd", "Juridische / technische aankoop-DD", 1000, "€", "Budgetreservering"),
        ("taxatie", "Taxatie, kadastrale & admin. kosten", 500, "€", "Budgetreservering"),
        ("verbouw_m2", "Verbouwbudget", p["focus_tier"], "€/m²", "Focus-tier uit DealRadar-profiel"),
        ("contingency", "Bouwcontingency / verborgen gebreken", 0.10, "%", "10% over directe verbouwkosten"),
        ("splitskosten", "Splitsingskosten (alleen bij splitsen)", p["split_kosten"], "€", "Leges, akte, VvE, advies"),
        ("rente", "Financieringsrente", p["rente_pct"] / 100, "%/jr", "Over de totale inleg"),
        ("looptijd", "Projectduur", p["looptijd_mnd"], "mnd", "Aankoop -> verkoop incl. verbouw"),
        ("holding_overig", "Overige holdingkosten", 0, "€", "Verzekering/OZB/nuts — vul aan per project"),
        ("courtage", "Verkoopcourtage", p["courtage_pct"] / 100, "% GDV", "Binnen NVM-indicatie ~1–2,5%"),
        ("verkoop_vast", "Vaste verkoopkosten", p["verkoop_vast"], "€", "Fotografie/brochure/notaris/royement"),
        ("split_verlies", "Verlies verkoopbaar opp. bij splitsen", p["go_verlies_pct"] / 100, "%", "Verkeersruimte na splitsing"),
        ("band", "Waardeband ± (conservatief scenario)", p["waarde_band"], "±", "Bandbreedte rond de exitwaarde"),
    ]
    for key, label, val, unit, toel in rows:
        r = A[key]
        _set(ws, f"A{r}", label, BLACK, border=True)
        fmt = PCT if unit in ("%", "%/jr", "% GDV", "±") else (EUR if unit == "€" or unit == "€/m²" else None)
        _set(ws, f"B{r}", val, BLUE, fmt=fmt, align="right", border=True)
        _set(ws, f"C{r}", unit, BLACK, border=True)
        _set(ws, f"D{r}", toel, Font(name=FONT, size=9, color="666666"), border=True)
    note = ws["A19"]
    note.value = "Blauwe cellen zijn aan te passen. De object-tabs verwijzen naar deze aannames en herrekenen automatisch."
    note.font = Font(name=FONT, italic=True, size=9, color="666666")


AN = "Aannames"   # sheet-naam voor referenties


def _ref(key: str) -> str:
    return f"'{AN}'!$B${A[key]}"


def _build_object_sheet(ws, o: dict, exit_m2: float, strategie: str, p: dict):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["D"].width = 26
    ws.column_dimensions["E"].width = 14

    addr = o.get("address") or "Object"
    plaats = o.get("city") or ""
    buurt = o.get("neighbourhood") or ""
    _set(ws, "A1", f"{addr} — DEVELOPMENT P&L", TITLE)
    _set(ws, "A2", f"{plaats}{' / ' + buurt if buurt else ''} · strategie: {strategie} · "
                   f"projectduur = Aannames", Font(name=FONT, size=9, color="666666"))

    # ── invoer (blauw) ──
    _set(ws, "A4", "OBJECT-INVOER", HEAD, fill=HEAD_FILL); _set(ws, "B4", "", HEAD, fill=HEAD_FILL)
    _set(ws, "A5", "Aankoopprijs (vraagprijs)")
    _set(ws, "B5", o.get("price") or 0, BLUE, fmt=EUR, align="right")
    _set(ws, "A6", "Woonoppervlak")
    _set(ws, "B6", o.get("living_area") or 0, BLUE, fmt='#,##0 "m²"', align="right")
    _set(ws, "A7", "Exit-verkoopprijs (gerenoveerd)")
    _set(ws, "B7", round(exit_m2), BLUE, fmt='€ #,##0 "/m²"', align="right")
    _set(ws, "A8", "Strategie")
    _set(ws, "B8", strategie, BLUE, align="right")
    _set(ws, "A9", "Verkoopbaar oppervlak")
    # bij splitsen verlies je verkeersruimte
    _set(ws, "B9", f'=IF(B8="splitsen",B6*(1-{_ref("split_verlies")}),B6)', BLACK, fmt='#,##0 "m²"', align="right")

    # ── opbrengsten ──
    _set(ws, "A11", "OPBRENGSTEN", HEAD, fill=HEAD_FILL); _set(ws, "B11", "", HEAD, fill=HEAD_FILL)
    _set(ws, "A12", "Bruto verkoopopbrengst (GDV)")
    _set(ws, "B12", "=B9*B7", BOLD, fmt=EUR, align="right")

    # ── aankoop & acquisitie ──
    _set(ws, "A14", "AANKOOP & ACQUISITIEKOSTEN", HEAD, fill=HEAD_FILL); _set(ws, "B14", "", HEAD, fill=HEAD_FILL)
    _set(ws, "A15", "Aankoopprijs"); _set(ws, "B15", "=B5", BLACK, fmt=EUR, align="right")
    _set(ws, "A16", "Overdrachtsbelasting"); _set(ws, "B16", f"=B5*{_ref('ovb')}", BLACK, fmt=EUR, align="right")
    _set(ws, "A17", "Notaris leveringsakte + Kadaster"); _set(ws, "B17", f"={_ref('notaris')}", BLACK, fmt=EUR, align="right")
    _set(ws, "A18", "Juridische / technische aankoop-DD"); _set(ws, "B18", f"={_ref('dd')}", BLACK, fmt=EUR, align="right")
    _set(ws, "A19", "Taxatie, kadastrale & admin. kosten"); _set(ws, "B19", f"={_ref('taxatie')}", BLACK, fmt=EUR, align="right")
    _set(ws, "A20", "Totaal acquisitie incl. aankoopprijs", BOLD); _set(ws, "B20", "=SUM(B15:B19)", BOLD, fmt=EUR, align="right")

    # ── verbouw / BOQ ──
    _set(ws, "A22", "VERBOUW (BOQ)", HEAD, fill=HEAD_FILL); _set(ws, "B22", "", HEAD, fill=HEAD_FILL)
    _set(ws, "A23", "Directe verbouwkosten"); _set(ws, "B23", f"=B6*{_ref('verbouw_m2')}", BLACK, fmt=EUR, align="right")
    _set(ws, "A24", "Bouwcontingency"); _set(ws, "B24", f"=B23*{_ref('contingency')}", BLACK, fmt=EUR, align="right")
    _set(ws, "A25", "Splitsingskosten"); _set(ws, "B25", f'=IF(B8="splitsen",{_ref("splitskosten")},0)', BLACK, fmt=EUR, align="right")
    _set(ws, "A26", "Totaal verbouw", BOLD); _set(ws, "B26", "=SUM(B23:B25)", BOLD, fmt=EUR, align="right")

    # ── investering + holding ──
    _set(ws, "A28", "INVESTERING & HOLDING", HEAD, fill=HEAD_FILL); _set(ws, "B28", "", HEAD, fill=HEAD_FILL)
    _set(ws, "A29", "Totale investering (acquisitie + verbouw)", BOLD); _set(ws, "B29", "=B20+B26", BOLD, fmt=EUR, align="right")
    _set(ws, "A30", "Financieringskosten (rente × looptijd)")
    _set(ws, "B30", f"=B29*{_ref('rente')}*{_ref('looptijd')}/12", BLACK, fmt=EUR, align="right")
    _set(ws, "A31", "Overige holdingkosten"); _set(ws, "B31", f"={_ref('holding_overig')}", BLACK, fmt=EUR, align="right")
    _set(ws, "A32", "Totaal holding", BOLD); _set(ws, "B32", "=B30+B31", BOLD, fmt=EUR, align="right")

    # ── verkoopkosten ──
    _set(ws, "A34", "VERKOOPKOSTEN", HEAD, fill=HEAD_FILL); _set(ws, "B34", "", HEAD, fill=HEAD_FILL)
    _set(ws, "A35", "Verkoopcourtage"); _set(ws, "B35", f"=B12*{_ref('courtage')}", BLACK, fmt=EUR, align="right")
    _set(ws, "A36", "Vaste verkoopkosten"); _set(ws, "B36", f"={_ref('verkoop_vast')}", BLACK, fmt=EUR, align="right")
    _set(ws, "A37", "Totaal verkoopkosten", BOLD); _set(ws, "B37", "=SUM(B35:B36)", BOLD, fmt=EUR, align="right")

    # ── resultaat ──
    _set(ws, "A39", "PROJECTRESULTAAT", HEAD, fill=HEAD_FILL); _set(ws, "B39", "", HEAD, fill=HEAD_FILL)
    _set(ws, "A40", "Totale projectkosten"); _set(ws, "B40", "=B29+B32+B37", BLACK, fmt=EUR, align="right")
    _set(ws, "A41", "PROJECTWINST (mid)", BOLD); _set(ws, "B41", "=B12-B40", BOLD, fmt=EUR, align="right", fill=SUB_FILL)
    _set(ws, "A42", "ROI op totale kosten", BOLD); _set(ws, "B42", "=B41/B40", BOLD, fmt=PCT, align="right", fill=SUB_FILL)
    _set(ws, "A43", "Geannualiseerde ROI"); _set(ws, "B43", f"=B42*12/{_ref('looptijd')}", BLACK, fmt=PCT, align="right")
    _set(ws, "A44", "Winstmarge op GDV"); _set(ws, "B44", "=B41/B12", BLACK, fmt=PCT, align="right")

    # ── conservatief scenario ──
    _set(ws, "A46", "CONSERVATIEF (exit − waardeband)", HEAD, fill=HEAD_FILL); _set(ws, "B46", "", HEAD, fill=HEAD_FILL)
    _set(ws, "A47", "GDV conservatief"); _set(ws, "B47", f"=B12*(1-{_ref('band')})", BLACK, fmt=EUR, align="right")
    _set(ws, "A48", "Projectwinst conservatief", BOLD); _set(ws, "B48", "=B47-B40", BOLD, fmt=EUR, align="right", fill=SUB_FILL)
    _set(ws, "A49", "ROI conservatief", BOLD); _set(ws, "B49", "=B48/B40", BOLD, fmt=PCT, align="right", fill=SUB_FILL)
    _set(ws, "A50", "Veiligheidsmarge (val vóór break-even)")
    _set(ws, "B50", "=B41/B12", BLACK, fmt=PCT, align="right")

    if o.get("url"):
        c = ws["A52"]; c.value = "Bekijk op Funda ↗"; c.hyperlink = o["url"]
        c.font = Font(name=FONT, color="1155CC", underline="single")


def build_pnl_workbook(out_dir: str, region: str = "grote_steden",
                       profile: str = "standaard", rank: str = "risk",
                       n: int = 10) -> str:
    """Genereer de Top-N P&L-workbook en geef het pad terug."""
    params = merged_params(get_profile(profile))
    data = top_listings(get_profile(profile), n=n, rank=rank, region=region)
    top = data["top"]
    medians = city_medians_all()

    wb = Workbook()
    ws_over = wb.active
    ws_over.title = "Top 10"
    ws_an = wb.create_sheet(AN)
    _build_aannames(ws_an, params)

    # ── overzichtstab ──
    ws_over.sheet_view.showGridLines = False
    headers = ["#", "Adres", "Stad", "Strategie", "Vraagprijs", "m²", "€/m²",
               "Exit €/m²", "Projectwinst", "ROI", "Marge", "Conservatief netto",
               "Deal-score", "Gevonden", "Funda"]
    widths = [4, 30, 16, 11, 13, 7, 9, 10, 14, 8, 8, 16, 11, 12, 8]
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        col = get_column_letter(i)
        ws_over.column_dimensions[col].width = w
        _set(ws_over, f"{col}3", h, HEAD, fill=HEAD_FILL,
             align="right" if i >= 5 else "left")
    _set(ws_over, "A1", "DEALRADAR — TOP 10 FLIP/DEVELOPMENT-KANSEN", TITLE)
    _set(ws_over, "A2",
         f"regio: {region} · profiel: {profile} · rang: {rank} · "
         f"gegenereerd {dt.datetime.now().strftime('%d-%m-%Y %H:%M')}",
         Font(name=FONT, size=9, color="666666"))

    used_names: set[str] = {AN, "Top 10"}
    row = 4
    for idx, o in enumerate(top, start=1):
        city_med = medians.get((o.get("city") or "").lower())
        if not city_med:
            continue
        strategie = "splitsen" if (o.get("best_strategie") == "splitsen"
                                   and o.get("split_allowed")) else "flip"
        huis_m2 = city_med * params["renov_uplift"]
        exit_m2 = huis_m2 * params["app_premium"] if strategie == "splitsen" else huis_m2

        sheet_name = _safe_sheet_name(f"{idx}. {o.get('address','')}", used_names)
        ws_obj = wb.create_sheet(sheet_name)
        _build_object_sheet(ws_obj, o, exit_m2, strategie, params)

        q = f"'{sheet_name}'"
        _set(ws_over, f"A{row}", idx, BLACK, border=True)
        _set(ws_over, f"B{row}", o.get("address") or "", BLACK, border=True)
        _set(ws_over, f"C{row}", o.get("city") or "", BLACK, border=True)
        _set(ws_over, f"D{row}", strategie, BLACK, border=True)
        _set(ws_over, f"E{row}", o.get("price") or 0, BLACK, fmt=EUR, align="right", border=True)
        _set(ws_over, f"F{row}", o.get("living_area") or 0, BLACK, fmt="#,##0", align="right", border=True)
        _set(ws_over, f"G{row}", o.get("price_m2") or 0, BLACK, fmt=EUR, align="right", border=True)
        _set(ws_over, f"H{row}", round(exit_m2), BLACK, fmt=EUR, align="right", border=True)
        # live links naar de object-tab
        _set(ws_over, f"I{row}", f"={q}!$B$41", GREEN, fmt=EUR, align="right", border=True)
        _set(ws_over, f"J{row}", f"={q}!$B$42", GREEN, fmt=PCT, align="right", border=True)
        _set(ws_over, f"K{row}", f"={q}!$B$44", GREEN, fmt=PCT, align="right", border=True)
        _set(ws_over, f"L{row}", f"={q}!$B$48", GREEN, fmt=EUR, align="right", border=True)
        _set(ws_over, f"M{row}", o.get("deal_score") or 0, BLACK, fmt="#,##0", align="right", border=True)
        # datum waarop de kans is gevonden
        fs = o.get("first_seen")
        if fs:
            try:
                _set(ws_over, f"N{row}", dt.datetime.fromisoformat(fs).date(),
                     BLACK, fmt="dd-mm-yyyy", align="right", border=True)
            except Exception:
                _set(ws_over, f"N{row}", fs[:10], BLACK, align="right", border=True)
        else:
            _set(ws_over, f"N{row}", "—", BLACK, align="right", border=True)
        if o.get("url"):
            c = ws_over[f"O{row}"]; c.value = "↗"; c.hyperlink = o["url"]
            c.font = Font(name=FONT, color="1155CC", underline="single")
            c.border = BORDER
        row += 1

    ws_over.freeze_panes = "A4"

    Path(out_dir).mkdir(parents=True, exist_ok=True)
    fname = f"DealRadar-Top{n}-PnL-{dt.date.today().isoformat()}.xlsx"
    path = str(Path(out_dir) / fname)
    wb.save(path)
    return path
